"""
ASMR Video Generation Workflow
Automated workflow for generating ASMR videos using Google Gemini and Veo 3
"""

import time
import os
import logging
from datetime import datetime
from google import genai
from google.genai import types
import openpyxl
from pathlib import Path
import re

# --- CONFIGURATION ---
os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = "oopsbot-8b65ad33386a.json"
PROJECT_ID = "oopsbot"
LOCATION = "us-central1"
EXCEL_FILE = "asmr_video_ideas.xlsx"
OUTPUT_FOLDER = "output"
LOGS_FOLDER = "logs"

# Create necessary folders
os.makedirs(OUTPUT_FOLDER, exist_ok=True)
os.makedirs(LOGS_FOLDER, exist_ok=True)

# --- LOGGING SETUP ---
log_filename = f"{LOGS_FOLDER}/workflow_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(log_filename, encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# --- INITIALIZE CLIENT ---
try:
    client = genai.Client(
        vertexai=True,
        project=PROJECT_ID,
        location=LOCATION
    )
    logger.info("✓ Successfully connected to Vertex AI")
except Exception as e:
    logger.error(f"✗ Failed to initialize Vertex AI client: {e}")
    exit(1)


def sanitize_filename(text):
    """Convert text to valid filename"""
    # Remove or replace invalid characters
    text = re.sub(r'[<>:"/\\|?*]', '', text)
    # Replace spaces with underscores
    text = text.replace(' ', '_')
    # Limit length
    return text[:50]


def load_or_create_excel():
    """Load existing Excel file or create new one"""
    if os.path.exists(EXCEL_FILE):
        logger.info(f"✓ Loading existing Excel file: {EXCEL_FILE}")
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        logger.info(f"✓ Creating new Excel file: {EXCEL_FILE}")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ASMR Ideas"
        
        # Create headers
        headers = ['Content Idea', 'Video Generation Prompt', 'Status', 'Video Filename']
        for col, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=header)
        
        # Set column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 80
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 40
        
        wb.save(EXCEL_FILE)
    
    return wb, ws


def generate_content_ideas(count=5):
    """Use Gemini to generate highly detailed ASMR content ideas"""
    logger.info(f"\n{'='*60}")
    logger.info("GENERATING HIGH-QUALITY CONTENT IDEAS")
    logger.info(f"{'='*60}")
    
    prompt = f"""Generate {count} UNIQUE and HIGHLY DETAILED ASMR video content ideas optimized for viral Instagram/TikTok content.

REQUIREMENTS for each idea:
- Be SPECIFIC about materials, actions, and visual elements
- Include details about colors, textures, and movements
- Mention camera angles or shot types that would work best
- Focus on proven ASMR triggers with high viral potential
- Ensure ideas are visually stunning and satisfying
- Make ideas detailed enough to guide professional video generation

TRENDING ASMR CATEGORIES (choose diverse types):
1. Slicing/Cutting (soap, kinetic sand, clay, foam, vegetables)
2. Pouring/Mixing (paint, resin, glitter, honey, colored liquids)
3. Organization (color-coordinated items, satisfying arrangements)
4. Textures (crumbling, squishing, peeling, unwrapping)
5. Creation (resin art, slime making, food styling, crafts)
6. Destruction (controlled demolition, pressure washing, cleaning)

QUALITY STANDARDS:
- Each idea should be 15-25 words
- Include specific colors, materials, and actions
- Mention visual style (macro, slow-motion, overhead, etc.)
- Think "premium commercial aesthetic"
- Viral potential is KEY

GOOD EXAMPLE:
"Extreme macro slow-motion shot of translucent teal soap being precision-sliced into perfect cubes with satisfying clean cuts, studio lighting revealing rainbow micro-bubbles"

BAD EXAMPLE:
"Soap cutting video"

Generate {count} DETAILED, VIRAL-READY ASMR video ideas now:"""
    
    try:
        logger.info("Requesting content ideas from Gemini...")
        response = client.models.generate_content(
            model='gemini-2.0-flash-exp',
            contents=prompt
        )
        
        ideas_text = response.text.strip()
        logger.info(f"✓ Received response from Gemini")
        logger.debug(f"Raw response:\n{ideas_text}")
        
        # Parse ideas (split by lines and clean)
        ideas = []
        for line in ideas_text.split('\n'):
            line = line.strip()
            # Remove numbering (1., 1), etc.)
            cleaned = re.sub(r'^\d+[\.\)]\s*', '', line)
            if cleaned and len(cleaned) > 10:  # Skip empty or very short lines
                ideas.append(cleaned)
        
        ideas = ideas[:count]  # Ensure we only get the requested count
        
        logger.info(f"✓ Generated {len(ideas)} content ideas:")
        for i, idea in enumerate(ideas, 1):
            logger.info(f"  {i}. {idea[:80]}...")
        
        return ideas
        
    except Exception as e:
        logger.error(f"✗ Error generating content ideas: {e}")
        logger.exception("Full exception details:")
        return []


def generate_video_prompt(content_idea):
    """Convert content idea into highly detailed, optimized video generation prompt"""
    logger.info(f"\nGenerating HIGH-QUALITY video prompt for: {content_idea[:50]}...")
    
    prompt = f"""You are a professional cinematographer and AI video generation expert specializing in creating prompts for Google Veo 3.

Content Idea: {content_idea}

Create an extremely detailed, professional video generation prompt that will produce MAXIMUM QUALITY output. Include:

REQUIRED ELEMENTS (ALL must be specified):
1. **Camera Work**: Exact camera movement (slow push-in, dolly zoom, crane shot, tracking shot, etc.)
2. **Shot Type**: Specific framing (extreme close-up, macro shot, wide angle, etc.)
3. **Lighting**: Detailed lighting setup (soft diffused lighting, dramatic side lighting, golden hour, studio lighting, etc.)
4. **Quality Markers**: 8K resolution, cinematic, RAW footage, professional grade, high dynamic range
5. **Visual Details**: Textures, colors, materials, reflections, particles - be VERY specific
6. **Composition**: Rule of thirds, depth of field, bokeh, focus points
7. **Motion**: Speed and type of movement (slow motion, time-lapse, smooth, dynamic)
8. **Atmosphere**: Mood, ambiance, visual effects, post-processing style

QUALITY STANDARDS:
- Think like a professional commercial director
- Use technical cinematography terms
- Specify exact visual qualities that make footage premium
- Include sensory details that enhance ASMR appeal
- Focus on what makes Instagram/TikTok content go viral

LENGTH: 250-400 characters (detailed but focused)

EXAMPLE STYLE:
"Extreme macro 8K shot of iridescent soap being precision-sliced, razor-sharp focus on crystalline edges, soft diffused studio lighting with rim highlights, slow-motion at 120fps capturing satisfying clean cuts, shallow depth of field with creamy bokeh, professional color grading with vibrant pastels, smooth motorized slider movement, ultra-detailed textures revealing micro-bubbles and rainbow reflections, cinematic composition, ASMR-optimized visual appeal"

Now create a professional prompt for: {content_idea}

Return ONLY the video generation prompt, nothing else."""
    
    try:
        response = client.models.generate_content(
            model='gemini-2.0-flash-exp',
            contents=prompt
        )
        
        video_prompt = response.text.strip()
        # Remove quotes if present
        video_prompt = video_prompt.strip('"\'')
        
        logger.info(f"✓ Generated HIGH-QUALITY prompt ({len(video_prompt)} chars)")
        logger.info(f"  Preview: {video_prompt[:100]}...")
        return video_prompt
        
    except Exception as e:
        logger.error(f"✗ Error generating video prompt: {e}")
        # Enhanced fallback with more detail
        return f"Cinematic 8K macro shot of {content_idea}, professional studio lighting with soft diffused key light, slow smooth camera movement, shallow depth of field, ultra-detailed textures, high dynamic range, professional color grading, ASMR-optimized visual clarity"


def generate_video(prompt, output_filename):
    """Generate video using Veo 3"""
    logger.info(f"\n{'='*60}")
    logger.info("GENERATING VIDEO")
    logger.info(f"{'='*60}")
    logger.info(f"Prompt: {prompt}")
    logger.info(f"Output: {output_filename}")
    
    try:
        # Generate video
        operation = client.models.generate_videos(
            model="veo-3.1-generate-preview",
            prompt=prompt,
            config=types.GenerateVideosConfig(
                aspect_ratio="9:16",  # Instagram/TikTok format
            )
        )
        
        logger.info("✓ Video generation request sent")
        logger.info(f"  Operation ID: {operation.name if hasattr(operation, 'name') else 'Unknown'}")
        logger.info("⏳ Waiting for video rendering (this may take 1-2 minutes)...")
        
        # Poll for completion
        poll_count = 0
        while not operation.done:
            time.sleep(10)
            poll_count += 1
            logger.info(f"  ⏳ Checking status... ({poll_count * 10}s elapsed)")
            operation = client.operations.get(operation)
        
        logger.info(f"✓ Video generation completed after {poll_count * 10} seconds")
        
        # Extract video
        if not operation.result or not operation.result.generated_videos:
            logger.error("✗ No video in result")
            return False
        
        video_result = operation.result.generated_videos[0]
        video_bytes = video_result.video.video_bytes if hasattr(video_result.video, 'video_bytes') else None
        
        if not video_bytes or len(video_bytes) == 0:
            logger.error("✗ No video bytes received")
            return False
        
        # Save video
        output_path = os.path.join(OUTPUT_FOLDER, output_filename)
        with open(output_path, 'wb') as f:
            f.write(video_bytes)
        
        file_size = os.path.getsize(output_path) / (1024 * 1024)
        logger.info(f"✓ Video saved successfully!")
        logger.info(f"  Location: {output_path}")
        logger.info(f"  Size: {file_size:.2f} MB")
        
        return True
        
    except Exception as e:
        logger.error(f"✗ Error generating video: {e}")
        logger.exception("Full exception details:")
        return False


def check_incomplete_ideas(ws):
    """Check for ideas that don't have videos generated yet"""
    incomplete = []
    
    for row_idx in range(2, ws.max_row + 1):
        idea = ws.cell(row=row_idx, column=1).value
        status = ws.cell(row=row_idx, column=3).value
        
        if idea and status != "Completed":
            incomplete.append(row_idx)
    
    return incomplete


def run_workflow():
    """Main workflow execution"""
    logger.info("\n" + "="*60)
    logger.info("ASMR VIDEO GENERATION WORKFLOW")
    logger.info(f"Started at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    logger.info("="*60)
    
    # Load Excel
    wb, ws = load_or_create_excel()
    
    # Check for incomplete ideas
    incomplete_rows = check_incomplete_ideas(ws)
    
    if incomplete_rows:
        logger.info(f"\n✓ Found {len(incomplete_rows)} incomplete idea(s)")
        logger.info("Continuing with incomplete ideas...")
    else:
        logger.info("\n✓ All previous ideas completed (or no ideas exist)")
        logger.info("Generating 5 new content ideas...")
        
        # Generate new ideas
        ideas = generate_content_ideas(count=5)
        
        if not ideas:
            logger.error("✗ Failed to generate ideas. Exiting.")
            return
        
        # Add ideas to Excel
        start_row = ws.max_row + 1
        for i, idea in enumerate(ideas):
            row = start_row + i
            ws.cell(row=row, column=1, value=idea)
            ws.cell(row=row, column=3, value="Pending")
            logger.info(f"  Added idea {i+1} to Excel at row {row}")
        
        wb.save(EXCEL_FILE)
        logger.info(f"✓ Excel updated with new ideas")
        
        # Update incomplete rows list
        incomplete_rows = check_incomplete_ideas(ws)
    
    # Process first incomplete idea
    if incomplete_rows:
        row_idx = incomplete_rows[0]
        idea = ws.cell(row=row_idx, column=1).value
        
        logger.info(f"\n{'='*60}")
        logger.info(f"PROCESSING IDEA #{row_idx-1}")
        logger.info(f"{'='*60}")
        logger.info(f"Idea: {idea}")
        
        # Check if prompt already exists
        existing_prompt = ws.cell(row=row_idx, column=2).value
        if existing_prompt:
            video_prompt = existing_prompt
            logger.info(f"✓ Using existing prompt from Excel")
        else:
            # Generate prompt
            video_prompt = generate_video_prompt(idea)
            ws.cell(row=row_idx, column=2, value=video_prompt)
            wb.save(EXCEL_FILE)
            logger.info(f"✓ Prompt saved to Excel")
        
        # Generate filename
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_idea = sanitize_filename(idea)
        filename = f"asmr_{row_idx-1}_{safe_idea}_{timestamp}.mp4"
        
        # Update status to "In Progress"
        ws.cell(row=row_idx, column=3, value="In Progress")
        wb.save(EXCEL_FILE)
        
        # Generate video
        success = generate_video(video_prompt, filename)
        
        if success:
            # Update Excel with completion
            ws.cell(row=row_idx, column=3, value="Completed")
            ws.cell(row=row_idx, column=4, value=filename)
            wb.save(EXCEL_FILE)
            logger.info(f"✓ Excel updated - Idea marked as Completed")
        else:
            ws.cell(row=row_idx, column=3, value="Failed")
            wb.save(EXCEL_FILE)
            logger.error(f"✗ Video generation failed - Idea marked as Failed")
    
    # Final summary
    logger.info(f"\n{'='*60}")
    logger.info("WORKFLOW SUMMARY")
    logger.info(f"{'='*60}")
    
    remaining = check_incomplete_ideas(ws)
    completed_count = ws.max_row - 1 - len(remaining)  # -1 for header row
    
    logger.info(f"Total ideas: {ws.max_row - 1}")
    logger.info(f"Completed: {completed_count}")
    logger.info(f"Remaining: {len(remaining)}")
    logger.info(f"Excel file: {EXCEL_FILE}")
    logger.info(f"Log file: {log_filename}")
    logger.info(f"\n✓ Workflow completed!")


if __name__ == "__main__":
    run_workflow()

