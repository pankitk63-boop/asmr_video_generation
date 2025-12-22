"""Quick script to display workflow status from Excel"""
import openpyxl
import os

if not os.path.exists('asmr_video_ideas.xlsx'):
    print("No Excel file found. Run workflow.py first!")
    exit(1)

wb = openpyxl.load_workbook('asmr_video_ideas.xlsx')
ws = wb.active

print("\n" + "="*80)
print("ASMR VIDEO GENERATION STATUS")
print("="*80)

for row_idx in range(1, ws.max_row + 1):
    idea = ws.cell(row_idx, 1).value
    prompt = ws.cell(row_idx, 2).value
    status = ws.cell(row_idx, 3).value
    filename = ws.cell(row_idx, 4).value
    
    if row_idx == 1:
        # Header
        print(f"\n{idea:45} | {prompt:45} | {status:12} | {filename if filename else 'N/A'}")
        print("-"*80)
    else:
        # Data rows
        idea_short = (idea[:42] + "...") if idea and len(idea) > 45 else (idea or "")
        prompt_short = (prompt[:42] + "...") if prompt and len(prompt) > 45 else (prompt or "N/A")
        filename_short = filename if filename else "N/A"
        print(f"{row_idx-1}. {idea_short:42} | {status:12} | {filename_short}")

# Summary
completed = sum(1 for i in range(2, ws.max_row + 1) if ws.cell(i, 3).value == "Completed")
total = ws.max_row - 1

print("\n" + "="*80)
print(f"Progress: {completed}/{total} completed ({(completed/total*100) if total > 0 else 0:.1f}%)")
print("="*80 + "\n")

